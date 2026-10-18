"""Export exact saved grid values, clipped to one server-resolved process."""
import io
import json
import numpy as np
from contracts import timestamp
from common import read_json

VARIABLES = {
    'WU': ('上层张力水', 'mm'), 'WL': ('下层张力水', 'mm'), 'WD': ('深层张力水', 'mm'),
    'S': ('自由水蓄量', 'mm'), 'FR': ('产流面积比例', '1'), 'RS': ('地表径流', 'mm'),
    'RI': ('壤中径流', 'mm'), 'RG': ('地下径流', 'mm'), 'W_WM': ('张力水相对蓄量', '1'),
    'precipitation': ('降雨', 'mm'), 'pet': ('潜在蒸散发', 'mm'),
    'et_actual': ('实际蒸散发', 'mm'), 'water_balance_error': ('水量平衡误差', 'mm')}


def process_history(manager, result_id, process_id):
    process = manager.results.process(result_id, process_id)
    from task_diagnostics import TaskDiagnostics
    if not hasattr(manager, 'diagnostics'):
        manager.diagnostics = TaskDiagnostics(manager)
    folder, metadata, history = manager.diagnostics._saved(process['sourceTaskId'])
    mask = np.array([timestamp(process['start']) <= timestamp(value) <= timestamp(process['end']) for value in metadata['times']])
    expected = int((timestamp(process['end']) - timestamp(process['start'])).total_seconds() / 3600) + 1
    if mask.sum() != expected:
        raise ValueError('逐时状态未覆盖当前洪水完整时段，请按原条件生成状态分布')
    fields = {key: history[key][mask] for key in metadata['fields']}
    fields['W_WM'] = (fields['WU'] + fields['WL'] + fields['WD']) / metadata['parameters']['WM']
    return process, folder, metadata, history['times'][mask].tolist(), fields


def process_ranges(manager, result_id, process_id):
    process, _, metadata, times, fields = process_history(manager, result_id, process_id)
    return {'processId': process['id'], 'start': times[0], 'end': times[-1],
            'ranges': {name: {'min': float(values.min()), 'max': float(values.max())} for name, values in fields.items()},
            'variables': [{'key': key, 'name': value[0], 'unit': value[1]} for key, value in VARIABLES.items() if key in fields],
            'parameters': metadata['parameters']}


def export_state(manager, result_id, body):
    if not isinstance(body, dict) or set(body) != {'processId', 'gridId', 'variables'}:
        raise ValueError('请选择洪水、网格和导出变量')
    process, folder, metadata, times, fields = process_history(manager, result_id, body['processId'])
    selected = body['variables']
    if not isinstance(selected, list) or not selected or any(not isinstance(key, str) for key in selected) or len(set(selected)) != len(selected) or any(key not in VARIABLES or key not in fields for key in selected):
        raise ValueError('导出变量未保存或选择无效')
    if body['gridId'] not in metadata['gridIds']:
        raise ValueError('当前计算没有该网格')
    index = metadata['gridIds'].index(body['gridId'])
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    book = Workbook()
    sheet = book.active
    sheet.title = '逐时数据'
    sheet.append(['时间', *[f'{VARIABLES[key][0]} {key} ({VARIABLES[key][1]})' for key in selected]])
    for step, time in enumerate(times):
        sheet.append([timestamp(time), *[float(fields[key][step, index]) for key in selected]])
        sheet.cell(sheet.max_row, 1).number_format = 'yyyy-mm-dd hh:mm'
    sheet.freeze_panes = 'B2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions['A'].width = 23
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='356D59')
        if cell.column > 1:
            sheet.column_dimensions[cell.column_letter].width = 28
    info = book.create_sheet('计算说明')
    request = read_json(folder / 'request.json')
    records = {'成果编号': result_id, '洪水过程': process['id'], '实际状态计算': process['sourceTaskId'],
               '开始时间': times[0], '结束时间': times[-1], '时间范围': '仅本场评价期，不含预热',
               '栅格': body['gridId'], '栅格位置': metadata['gridGeometry'][index],
               '参数': metadata['parameters'], '变量与单位': {key: VARIABLES[key] for key in selected},
               '派生公式': 'W_WM = (WU + WL + WD) / WM', '资料及程序版本': metadata['provenance'],
               '起算条件': request.get('initialization') or request.get('period'),
               '数值说明': '单元格使用已保存原始数值；显示格式不改变数值。时间为本地无时区整点。'}
    for key, value in records.items():
        # All complex/user-facing strings are literal text, never formulas.
        info.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list, tuple)) else str(value)])
        info.cell(info.max_row, 2).data_type = 's'
    info.column_dimensions['A'].width = 24
    info.column_dimensions['B'].width = 110
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
