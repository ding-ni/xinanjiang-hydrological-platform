import argparse
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DEFAULT_MD_PATH = ROOT / "模块化新安江水文模型可视化平台建设方案.md"
DEFAULT_DOCX_PATH = ROOT / "模块化新安江水文模型可视化平台建设方案.docx"
DEFAULT_TITLE = "新安江水文模型业务化预报与可视化平台建设方案"
DEFAULT_SUBTITLE = "面向浙江试点流域的分区产汇流模拟、实时站网资料接入与预报研判体系"
DEFAULT_SCOPE = "浙江试点流域新安江模型业务化预报平台"
DEFAULT_DATE = "2026年6月"
DEFAULT_TAGLINE = "资料、模型、状态、数据库、NetCDF 与可视化的一体化方案"

BODY_FONT = "SimSun"
HEADING_FONT = "Microsoft YaHei"
ASCII_FONT = "Calibri"
CODE_FONT = "Consolas"

BLUE = RGBColor(31, 77, 120)
MID_BLUE = RGBColor(46, 116, 181)
NAVY = RGBColor(11, 37, 69)
MUTED = RGBColor(89, 89, 89)
LIGHT_FILL = "F4F6F9"
TABLE_BORDER = "B8C4D2"


def set_font(run, font=BODY_FONT, size=None, color=None, bold=None, italic=None):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ASCII_FONT if font == BODY_FONT else font)
    rfonts.set(qn("w:hAnsi"), ASCII_FONT if font == BODY_FONT else font)
    rfonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, font=BODY_FONT, size=11, color=None, bold=None):
    style.font.name = font
    style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ASCII_FONT if font == BODY_FONT else font)
    rfonts.set(qn("w:hAnsi"), ASCII_FONT if font == BODY_FONT else font)
    rfonts.set(qn("w:eastAsia"), font)


def set_paragraph_spacing(paragraph, before=0, after=8, line=1.333):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for edge, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=TABLE_BORDER):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_width(table, col_widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    width_total = sum(col_widths)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in col_widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(col_widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(col_widths[idx]))
            tc_w.set(qn("w:type"), "dxa")


def mark_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def mark_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def add_field(run, instruction):
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_header_footer(section, title):
    section.different_first_page_header_footer = True
    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header_p.add_run(title)
    set_font(r, font=BODY_FONT, size=9, color=MUTED)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = footer_p.add_run("第 ")
    set_font(r1, font=BODY_FONT, size=9, color=MUTED)
    r2 = footer_p.add_run()
    set_font(r2, font=BODY_FONT, size=9, color=MUTED)
    add_field(r2, "PAGE")
    r3 = footer_p.add_run(" 页")
    set_font(r3, font=BODY_FONT, size=9, color=MUTED)


def setup_document(title):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    add_header_footer(section, title)

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, BODY_FONT, 11, RGBColor(0, 0, 0))
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for style_name, size, color, before, after in (
        ("Heading 1", 16, MID_BLUE, 18, 10),
        ("Heading 2", 13, MID_BLUE, 12, 6),
        ("Heading 3", 12, BLUE, 8, 4),
    ):
        style = styles[style_name]
        set_style_font(style, HEADING_FONT, size, color, True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        set_style_font(style, BODY_FONT, 11, RGBColor(0, 0, 0))
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", 1)
    else:
        code_style = styles["Code Block"]
    set_style_font(code_style, CODE_FONT, 9, RGBColor(50, 50, 50))
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(8)
    code_style.paragraph_format.line_spacing = 1.0

    return doc


def add_cover(doc, title, subtitle, scope, date, tagline):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(92)
    spacer.paragraph_format.space_after = Pt(0)

    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p1.add_run(title)
    set_font(r, HEADING_FONT, 24, NAVY, True)
    p1.paragraph_format.space_after = Pt(8)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("业务化预报与可视化平台建设方案")
    set_font(r, HEADING_FONT, 27, NAVY, True)
    p2.paragraph_format.space_after = Pt(10)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p3.add_run(subtitle)
    set_font(r, BODY_FONT, 13, MUTED)
    p3.paragraph_format.space_after = Pt(42)

    meta = doc.add_table(rows=3, cols=2)
    meta.style = "Table Grid"
    set_table_width(meta, [1900, 4300])
    set_table_borders(meta, "D7DEE8")
    rows = [
        ("文档类型", "建设方案"),
        ("适用范围", scope),
        ("编制日期", date),
    ]
    for row_idx, row_data in enumerate(rows):
        for col_idx, text in enumerate(row_data):
            cell = meta.cell(row_idx, col_idx)
            set_cell_margins(cell, 100, 100, 140, 140)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if col_idx == 0:
                set_cell_shading(cell, LIGHT_FILL)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            set_font(run, BODY_FONT, 10.5, NAVY if col_idx == 0 else RGBColor(0, 0, 0), col_idx == 0)

    tail = doc.add_paragraph()
    tail.paragraph_format.space_before = Pt(130)
    tail.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tail.add_run(tagline)
    set_font(r, BODY_FONT, 10.5, MUTED, False)
    tail.add_run().add_break(WD_BREAK.PAGE)


def parse_markdown(md_text):
    lines = md_text.splitlines()
    body_start = 0
    for i, line in enumerate(lines):
        if line.startswith("## 一、"):
            body_start = i
            break
    return lines[body_start:]


def collect_headings(lines):
    headings = []
    for line in lines:
        if line.startswith("## "):
            headings.append((1, line[3:].strip()))
        elif line.startswith("### "):
            headings.append((2, line[4:].strip()))
    return headings


def add_contents(doc, headings):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("目录")
    set_font(r, HEADING_FONT, 18, NAVY, True)
    p.paragraph_format.space_after = Pt(18)

    for level, title in headings:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25 if level == 2 else 0)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(title)
        set_font(r, BODY_FONT, 10.5, RGBColor(0, 0, 0), False)
    doc.add_page_break()


def split_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def is_separator_row(line):
    cells = split_table_row(line)
    return all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells)


def column_widths(ncols):
    if ncols == 2:
        return [2500, 6860]
    if ncols == 3:
        return [1950, 3600, 3810]
    if ncols == 4:
        return [1700, 1750, 3150, 2760]
    return [int(9360 / ncols)] * ncols


def add_markdown_table(doc, rows):
    header = rows[0]
    body = rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    set_table_width(table, column_widths(len(header)))
    set_table_borders(table)
    mark_repeat_header(table.rows[0])
    mark_cant_split(table.rows[0])

    for idx, text in enumerate(header):
        cell = table.cell(0, idx)
        set_cell_shading(cell, LIGHT_FILL)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_font(r, BODY_FONT, 9.5, NAVY, True)

    for row_data in body:
        row = table.add_row()
        mark_cant_split(row)
        for idx, text in enumerate(row_data):
            cell = row.cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if len(text) > 12 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            set_font(r, BODY_FONT, 9.3, RGBColor(0, 0, 0), False)

    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(4)


def add_body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, 0, 8, 1.333)
    r = p.add_run(text)
    set_font(r, BODY_FONT, 11, RGBColor(0, 0, 0))


def add_code_block(doc, lines):
    p = doc.add_paragraph(style="Code Block")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("\n".join(lines))
    set_font(r, CODE_FONT, 9, RGBColor(50, 50, 50))
    p_pr = p._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F7F8FA")
    p_pr.append(shd)


def render_body(doc, lines):
    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i].rstrip()

        if line.strip().startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not line.strip() or line.strip() == "---":
            i += 1
            continue

        if line.startswith("|") and i + 1 < len(lines) and is_separator_row(lines[i + 1]):
            table_rows = [split_table_row(line)]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                table_rows.append(split_table_row(lines[i]))
                i += 1
            add_markdown_table(doc, table_rows)
            continue

        if line.startswith("## "):
            p = doc.add_paragraph(line[3:].strip(), style="Heading 1")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            i += 1
            continue

        if line.startswith("### "):
            p = doc.add_paragraph(line[4:].strip(), style="Heading 2")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            i += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            r = p.add_run(numbered.group(1))
            set_font(r, BODY_FONT, 11, RGBColor(0, 0, 0))
            i += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(bullet.group(1))
            set_font(r, BODY_FONT, 11, RGBColor(0, 0, 0))
            i += 1
            continue

        add_body_paragraph(doc, line)
        i += 1


def main():
    parser = argparse.ArgumentParser(description="Render a Markdown construction plan as a styled Word document.")
    parser.add_argument("--input", type=Path, default=DEFAULT_MD_PATH)
    parser.add_argument("--output", type=Path, default=DEFAULT_DOCX_PATH)
    parser.add_argument("--title", default=DEFAULT_TITLE)
    parser.add_argument("--subtitle", default=DEFAULT_SUBTITLE)
    parser.add_argument("--scope", default=DEFAULT_SCOPE)
    parser.add_argument("--date", default=DEFAULT_DATE)
    parser.add_argument("--tagline", default=DEFAULT_TAGLINE)
    args = parser.parse_args()

    md_text = args.input.read_text(encoding="utf-8")
    lines = parse_markdown(md_text)
    headings = collect_headings(lines)

    doc = setup_document(args.title)
    add_cover(doc, args.title, args.subtitle, args.scope, args.date, args.tagline)
    add_contents(doc, headings)
    render_body(doc, lines)
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
