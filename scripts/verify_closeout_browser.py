"""End-to-end closeout acceptance against an isolated local workspace."""
import argparse
import io
import json
import time
from pathlib import Path
from playwright.sync_api import sync_playwright, expect
from openpyxl import load_workbook

BASE='OP-ba24bb27ac6d45d89248aea49f24d39a'
EVENT='xf-31200701'
SCHEME='SCHEME-ad84e50a028349cf9622637524216398'
parser=argparse.ArgumentParser()
parser.add_argument('--url',default='http://127.0.0.1:18918')
parser.add_argument('--output',type=Path,default=Path('开发产物/收尾复核20261018/browser-check'))
args=parser.parse_args();args.output.mkdir(parents=True,exist_ok=True)
report={'checks':[],'pageErrors':[],'screenshots':[]}
expect.set_options(timeout=45000)
with sync_playwright() as p:
    browser=p.chromium.launch(headless=True)
    context=browser.new_context(viewport={'width':1366,'height':768},accept_downloads=True)
    page=context.new_page()
    page.on('pageerror',lambda error:report['pageErrors'].append(str(error)))
    page.on('response',lambda response:report['pageErrors'].append('读取空方案编号：'+response.url) if '/model-schemes/null' in response.url else None)
    def api(path,body=None):
        response=context.request.get(args.url+'/api/model/'+path,timeout=90000) if body is None else context.request.post(args.url+'/api/model/'+path,data=body,timeout=90000)
        if not response.ok:raise RuntimeError(str(response.status)+' '+response.text())
        return response.json()
    def screenshot(label):
        name=f'{label}-{page.viewport_size["width"]}.png'
        page.screenshot(path=str(args.output/name),full_page=True)
        report['screenshots'].append(name)
    try:
        page.goto(args.url+'/#/state?operation='+BASE+'&event='+EVENT)
        page.wait_for_load_state('networkidle')
        expect(page.get_by_role('heading',name='当前时刻空间分布')).to_be_visible()
        expect(page.get_by_role('button',name='导出 Excel',exact=True)).to_be_enabled()
        expect(page.locator('.grid-id')).to_have_text('NX-001')
        process=next(row for row in api('result-catalog/'+BASE+'/processes')['processes'] if row['id']==EVENT)
        metadata=api('tasks/'+process['sourceTaskId']+'/grid-state/catalog')
        grid=next(row for row in metadata['gridGeometry'] if row['id']=='NX-002')
        vertices=grid['bounds'][:4]
        center=[sum(point[axis] for point in vertices)/len(vertices) for axis in [0,1]]
        canvas=page.locator('.map-canvas').first
        canvas.scroll_into_view_if_needed()
        page.wait_for_function('()=>{const map=document.querySelector(".map-canvas").__hydroMap;return map && map.loaded() && !map.isMoving()}')
        pixel=canvas.evaluate('(element,point)=>{const p=element.__hydroMap.project(point);return {x:p.x,y:p.y}}',center)
        box=canvas.bounding_box();page.mouse.click(box['x']+pixel['x'],box['y']+pixel['y'])
        expect(page.locator('.grid-id')).to_have_text('NX-002')
        range_select=page.get_by_text('色标范围',exact=False).locator('select').first
        range_select.select_option('fixed');range_select.select_option('event')
        page.get_by_role('button',name='上层张力水',exact=False).click()
        page.get_by_role('button',name='导出 Excel',exact=True).click()
        dialog=page.get_by_role('dialog').filter(has_text='导出选中网格状态')
        dialog.get_by_role('button',name='全选',exact=True).click()
        with page.expect_download() as download:
            dialog.get_by_role('button',name='下载 Excel',exact=True).click()
        path=args.output/'selected-grid.xlsx';download.value.save_as(path)
        book=load_workbook(path,data_only=True)
        assert book.sheetnames==['逐时数据','计算说明']
        assert book['逐时数据'].max_row==290
        assert book['逐时数据']['A2'].value.isoformat()=='2020-07-01T08:00:00'
        assert book['逐时数据']['A290'].value.isoformat()=='2020-07-13T08:00:00'
        exact=api('tasks/'+process['sourceTaskId']+'/grid-state?unitId=NX-002')
        index=exact['times'].index('2020-07-01T08:00:00')
        assert abs(book['逐时数据']['B2'].value-exact['series']['WU'][index])<1e-12
        report['checks'].append('selected grid multi-variable Excel clipped to exact flood')
        for width,height in [(1366,768),(1920,1080)]:
            page.set_viewport_size({'width':width,'height':height});screenshot('state')
            assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth+2')
        page.get_by_role('button',name='断面流量',exact=True).click()
        expect(page.get_by_role('heading',name='选择查询断面')).to_be_visible()
        select=page.get_by_label('已保存断面')
        expect(select).to_have_value('SECTION-XUFAN')
        assert select.get_attribute('size') is None
        report['checks'].append('ordinary single-line default Xufan section dropdown')
        screenshot('sections')
        page.goto(args.url+'/#/calibration?schemeId='+SCHEME)
        page.wait_for_load_state('networkidle')
        expect(page.get_by_role('button',name='编辑方案',exact=True)).to_be_visible()
        page.get_by_role('button',name='编辑方案',exact=True).click()
        page.get_by_role('button',name='连续时段',exact=True).click()
        expect(page.get_by_role('heading',name='评价时段',exact=True)).to_be_visible()
        page.get_by_role('button',name='洪水场次',exact=True).click()
        page.get_by_role('button',name='率定设置',exact=True).click()
        page.locator('#algorithm').select_option('SCE_UA')
        expect(page.get_by_text('实际种群',exact=False)).to_be_visible()
        screenshot('sce-settings')
        page.get_by_role('button',name='完成',exact=True).click()
        expect(page.get_by_role('dialog').filter(has_text='率定设置')).not_to_be_visible()
        for width,height in [(1366,768),(1920,1080)]:
            page.set_viewport_size({'width':width,'height':height});screenshot('calibration')
            assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth+2')
        report['checks'].append('continuous editor reachable, return to event, SCE-UA settings')
        page.get_by_role('button',name='放弃修改',exact=True).click()
        # Real temporary simulation exercises the actual submission and saved result UI.
        simulation=api('operations',{'operation':'eventSimulation','parameterSourceResultId':BASE,
            'events':[{'id':EVENT,'start':'2020-07-01T08:00','end':'2020-07-13T08:00'}]})
        identifier=simulation['taskId'];report['temporaryOperation']=identifier
        deadline=time.monotonic()+180
        while time.monotonic()<deadline:
            status=api('operations/'+identifier)
            if status['status'] not in ['running','queued','cancelling']:break
            time.sleep(.5)
        assert status['status']=='completed',status.get('errorMessage')
        page.goto(args.url+'/#/simulation?operation='+identifier)
        page.wait_for_load_state('networkidle')
        expect(page.get_by_role('button',name='保存成果',exact=True)).to_be_visible()
        expect(page.get_by_role('button',name='☆ 收藏',exact=True)).to_be_disabled()
        parameter=api('parameter-schemes',{'name':'隔离验收：仅保存参数','sourceTaskId':identifier})
        assert api('result-retention/'+identifier)['retention']=='temporary'
        page.get_by_role('button',name='状态分布',exact=True).click()
        expect(page.get_by_role('heading',name='当前时刻空间分布')).to_be_visible()
        assert not page.get_by_role('dialog').is_visible()
        page.get_by_role('link',name='计算成果',exact=True).first.click()
        leave=page.get_by_role('dialog').filter(has_text='处理本轮临时结果')
        expect(leave).to_be_visible();leave.locator('select').select_option('discard')
        leave.get_by_role('button',name='取消离开',exact=True).click()
        expect(page.get_by_role('heading',name='当前时刻空间分布')).to_be_visible()
        page.reload();page.wait_for_load_state('networkidle')
        page.get_by_role('link',name='计算成果',exact=True).first.click()
        leave=page.get_by_role('dialog').filter(has_text='处理本轮临时结果')
        expect(leave).to_be_visible();leave.locator('select').select_option('discard')
        leave.get_by_role('button',name='确认并离开',exact=True).click()
        expect(leave).not_to_be_visible()
        page.wait_for_load_state('networkidle')
        assert api('result-retention/'+identifier)['retention']=='trashed'
        assert identifier not in [row['id'] for row in api('result-catalog')['results']]
        assert api('parameter-schemes/'+parameter['id'])['parameters']==parameter['parameters']
        report['checks'].append('temporary simulation, internal state navigation, cancel leaving, reload recovery, discard on leaving')
        # Restore/save must be independent of favorite and preserve numerical data.
        state=api('result-retention/'+identifier)
        state=api('result-retention/'+identifier+'/restore',{'revision':state['revision']})
        state=api('result-retention/'+identifier+'/save',{'revision':state['revision']})
        assert state['retention']=='saved'
        assert not api('result-metadata/'+identifier).get('favorite')
        report['checks'].append('restore then save is independent of favorite')
        page.goto(args.url+'/#/calibration')
        page.wait_for_load_state('networkidle')
        page.get_by_role('button',name='＋ 新建方案',exact=True).click()
        page.locator('#new-scheme-name').fill('回收站验收方案')
        page.get_by_label('新方案计算方式').select_option('continuous')
        page.get_by_role('button',name='建立方案',exact=True).click()
        expect(page.get_by_role('heading',name='回收站验收方案',exact=True)).to_be_visible()
        page.get_by_role('button',name='删除方案',exact=True).click()
        deletion=page.get_by_role('dialog').filter(has_text='删除方案 ·')
        expect(deletion).to_be_visible();deletion.get_by_role('button',name='确认移入回收站',exact=True).click()
        expect(page.get_by_role('heading',name='我的率定方案')).to_be_visible()
        page.get_by_role('button',name='回收站',exact=True).click()
        page.get_by_role('button',name='恢复',exact=True).last.click()
        page.get_by_role('dialog').get_by_role('button',name='恢复方案',exact=True).click()
        report['checks'].append('create continuous scheme, preview deletion, trash, restore')
        assert not report['pageErrors'],report['pageErrors']
        report['passed']=True
    except Exception as error:
        report['error']=str(error);screenshot('failure')
        raise
    finally:
        (args.output/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
        browser.close()
