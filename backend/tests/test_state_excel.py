import io
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from test_input_http import InputHttpTests
from state_export import export_state, process_ranges


class StateExcelTests(InputHttpTests):
    def test_exact_selected_grid_and_derived_value(self):
        folder, meta=self.saved_history()
        process={'id':'flood','sourceTaskId':'JOB-history','start':meta['times'][1],'end':meta['times'][1]}
        with patch.object(self.manager.results,'process',return_value=process):
            data=export_state(self.manager,'JOB-history',{'processId':'flood','gridId':'NX-002','variables':['WL','W_WM']})
            book=load_workbook(io.BytesIO(data),data_only=True)
            rows=list(book['逐时数据'].values)
            self.assertEqual(len(rows),2)
            self.assertEqual(rows[1][0].isoformat(),meta['times'][1])
            self.assertEqual(rows[1][1],4.)
            self.assertEqual(rows[1][2],.12)
            self.assertEqual(book.sheetnames,['逐时数据','计算说明'])
            ranges=process_ranges(self.manager,'JOB-history','flood')
            self.assertEqual(ranges['ranges']['WL'],{'min':3.,'max':4.})
            for body in [dict(processId='flood',gridId='bad',variables=['WL']),dict(processId='flood',gridId='NX-001',variables=['unknown'])]:
                with self.assertRaises(ValueError):export_state(self.manager,'JOB-history',body)
        process['end']='2017-06-19T13:00:00'
        with patch.object(self.manager.results,'process',return_value=process),self.assertRaisesRegex(ValueError,'完整时段'):
            process_ranges(self.manager,'JOB-history','flood')
